"""Excel output for shared accounting reports."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def profit_and_loss_xlsx(report: dict, start: date, end: date) -> bytes:
    """Build a one-sheet GBP P&L workbook from the shared report result."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Profit and loss"
    sheet.append(["Profit and loss (GBP)"])
    sheet.append(["From", start, "To", end])
    sheet.append(["Positive amounts are profit; negative amounts are loss."])
    sheet.append([])
    sheet.append(["Code", "Account", "Type", "Profit / loss GBP"])

    for row in report["rows"]:
        sheet.append([row["code"], row["name"], row["account_type"], float(row["profit"])])

    total_row = sheet.max_row + 1
    sheet.cell(total_row, 2, "Net profit / loss")
    sheet.cell(total_row, 4, float(report["profit"]))

    for row in (1, 5, total_row):
        for cell in sheet[row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234E70")
    sheet.column_dimensions["A"].width = 13
    sheet.column_dimensions["B"].width = 31
    sheet.column_dimensions["C"].width = 17
    sheet.column_dimensions["D"].width = 21
    sheet["B2"].number_format = "yyyy-mm-dd"
    sheet["D2"].number_format = "yyyy-mm-dd"
    for row in sheet.iter_rows(min_row=6, max_row=total_row, min_col=4, max_col=4):
        row[0].number_format = '#,##0.00;[Red](#,##0.00)'
    sheet.freeze_panes = "A6"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
